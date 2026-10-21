import io
import json
import os
import re
import threading
import time
from datetime import timedelta
from functools import wraps

import pandas as pd
import pyotp
import segno
from dotenv import load_dotenv
from flask import Flask, abort, jsonify, request, send_file, session
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select, update
from werkzeug.exceptions import HTTPException
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import check_password_hash, generate_password_hash

import analysis
import db
import security

VARS = {"rain_mm": "Rainfall (mm/month)", "t2m": "Mean temperature (°C)",
        "t2m_max": "Max temperature (°C)", "rh2m": "Relative humidity (%)",
        "era5_t2m": "ERA5 temperature (°C)", "era5_dewpoint": "ERA5 dew point (°C)",
        "era5_rain_mm": "ERA5 rainfall (mm/month)", "era5_rh2m": "ERA5 relative humidity (%)",
        "era5_radiation": "ERA5 solar radiation (W/m²)", "era5_wind": "ERA5 wind speed (m/s)",
        "era5_pressure": "ERA5 surface pressure (hPa)", "era5_soil": "ERA5 soil moisture (m³/m³)"}
SOURCES = [("rain_mm", "CHIRPS v2.0 (UCSB Climate Hazards Center), mean of 0.05° cells per LGA"),
           ("t2m, t2m_max, rh2m", "NASA POWER monthly, point at LGA centroid (~0.5° resolution, modelled)"),
           ("era5_* columns", "ERA5-Land monthly means (ECMWF / Copernicus C3S), ~9 km grid, 1996 to date, mean of grid cells "
                             "inside each LGA. Modelled reanalysis; ERA5 humidity is derived from mean temperature and dew point."),
           ("Survey indicators", "DHS (state level from the DHS API; LGA level from NDHS clusters), Malaria Atlas Project and "
                                 "IHME rasters (LGA zonal means). See the note column for sample size and caveats."),
           ("Note", "Modelled/satellite estimates, not station measurements. Validate against NiMet data where possible.")]
PERIOD = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
DUMMY_HASH = generate_password_hash("not-a-real-password")  # keeps login timing equal for unknown emails


def safe_cell(v):
    """Neutralise Excel formula injection (=, +, -, @ at the start of a text cell)."""
    return "'" + v if isinstance(v, str) and v and v[0] in "=+-@\t\r" else v


def nulls(df):
    return df.astype(object).where(df.notna(), None)


def create_app(engine=None):
    load_dotenv()
    prod = os.environ.get("APP_ENV") == "production"
    if prod and not os.environ.get("SECRET_KEY"):
        raise RuntimeError("SECRET_KEY must be set in production")
    engine = engine or db.get_engine()
    allowed_origin = os.environ.get("ALLOWED_ORIGIN", "http://localhost:3000")
    if prod and (not allowed_origin.startswith("https://") or "localhost" in allowed_origin or allowed_origin.endswith("/")):
        raise RuntimeError("In production set ALLOWED_ORIGIN to the frontend's https address with no trailing slash, "
                           "e.g. https://ondo-climate-web.onrender.com")

    app = Flask(__name__)
    app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Strict", SESSION_COOKIE_SECURE=prod,
                      PERMANENT_SESSION_LIFETIME=timedelta(hours=8), MAX_CONTENT_LENGTH=16 * 1024)
    app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY") or os.urandom(32).hex()
    vault = security.fernet(app.config["SECRET_KEY"])
    if prod:
        hops = int(os.environ.get("TRUSTED_PROXY_HOPS", "1"))  # verify with scripts/check_deploy.sh
        app.wsgi_app = ProxyFix(app.wsgi_app, x_for=hops)
    limiter = Limiter(get_remote_address, app=app, default_limits=["200/minute"],
                      storage_uri=os.environ.get("RATELIMIT_STORAGE_URI", "memory://"))

    @app.errorhandler(HTTPException)
    def http_error(e):
        msg = "Too many attempts. Try again later." if e.code == 429 else e.description
        return jsonify(error=msg), e.code

    @app.before_request
    def origin_guard():  # CSRF defence on top of SameSite=Strict cookies
        if request.method in ("POST", "PUT", "PATCH", "DELETE") and request.headers.get("Origin") != allowed_origin:
            abort(403, "Origin not allowed")

    @app.after_request
    def headers(r):
        r.headers["X-Content-Type-Options"] = "nosniff"
        r.headers["Referrer-Policy"] = "no-referrer"
        r.headers["Cache-Control"] = "no-store"
        r.headers["Content-Security-Policy"] = "default-src 'none'; frame-ancestors 'none'"
        if prod:
            r.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
        return r

    def user_by_id(uid):
        with engine.connect() as c:
            return c.execute(select(db.users).where(db.users.c.id == uid)).first()

    def email_key():
        return str((request.get_json(silent=True) or {}).get("email", "")).lower()[:254]

    def start_session(user):
        session.clear()
        session["uid"], session["pv"] = user.id, user.password_hash[-16:]  # pv: a password change ends old sessions
        session.permanent = True
        return jsonify(email=user.email)

    def login_required(fn):
        @wraps(fn)
        def wrapper(*a, **kw):
            user = user_by_id(session.get("uid"))
            if not user or session.get("pv") != user.password_hash[-16:]:
                session.pop("uid", None)  # keep pending_uid: a stray request must not end a half-finished 2FA sign-in
                session.pop("pv", None)
                abort(401, "Authentication required")
            return fn(*a, **kw)
        return wrapper

    def fetch(by_lga=True, start=None, end=None):
        """Validated, parameterised query. Every user input is allow-listed or pattern-checked, then bound."""
        names = request.args.getlist("lga")
        start, end = start or request.args.get("start"), end or request.args.get("end")
        for p in (start, end):
            if p and not PERIOD.match(p):
                abort(400, "Dates must look like 2024-03")
        with engine.connect() as c:
            known = {r[0] for r in c.execute(select(db.lgas.c.name))}
            if len(names) > 50 or any(n not in known for n in names):
                abort(400, "Unknown LGA")
            t = db.climate.c
            q = select(db.climate).order_by(t.lga, t.year, t.month)
            if names and by_lga:
                q = q.where(t.lga.in_(names))
            if start:
                q = q.where(t.year * 100 + t.month >= int(start.replace("-", "")))
            if end:
                q = q.where(t.year * 100 + t.month <= int(end.replace("-", "")))
            df = pd.read_sql(q, c)
        df["period"] = df.year.astype(str) + "-" + df.month.astype(str).str.zfill(2)
        return df

    def variable():
        v = request.args.get("variable", "rain_mm")
        if v not in VARS:
            abort(400, "Unknown variable")
        return v

    @app.get("/api/health")
    def health():
        return jsonify(ok=True, ip=request.remote_addr)  # lets scripts/check_deploy.sh verify proxy settings

    @app.post("/api/auth/login")
    @limiter.limit("5/minute;30/hour")
    @limiter.limit("10/hour", key_func=email_key)
    def login():
        body = request.get_json(silent=True) or {}
        email = str(body.get("email", "")).strip().lower()[:254]
        password = str(body.get("password", ""))[:256]
        with engine.connect() as c:
            user = c.execute(select(db.users).where(db.users.c.email == email)).first()
        ok = check_password_hash(user.password_hash if user else DUMMY_HASH, password)
        if not (user and ok):
            abort(401, "Invalid email or password")
        session.clear()
        if user.totp_enabled:  # password was right; the second factor is still required
            session["pending_uid"], session["pending_at"] = user.id, time.time()
            return jsonify(mfa_required=True)
        return start_session(user)

    @app.post("/api/auth/mfa")
    @limiter.limit("5/minute;30/hour")
    @limiter.limit("10/hour", key_func=lambda: str(session.get("pending_uid", "")))
    def mfa_login():
        user = user_by_id(session.get("pending_uid")) if time.time() - session.get("pending_at", 0) < 300 else None
        code = str((request.get_json(silent=True) or {}).get("code", ""))
        if not (user and user.totp_enabled and security.verify_totp(vault.decrypt(user.totp_secret.encode()).decode(), code)):
            abort(401, "Invalid or expired code. Sign in again if this keeps happening.")
        return start_session(user)

    @app.post("/api/auth/mfa/setup")
    @login_required
    def mfa_setup():
        user = user_by_id(session["uid"])
        if user.totp_enabled:
            abort(409, "Two-factor sign-in is already on")
        secret = pyotp.random_base32()
        session["mfa_setup"] = secret
        uri = pyotp.TOTP(secret).provisioning_uri(name=user.email, issuer_name="Ondo Climate Data")
        return jsonify(secret=secret, qr=segno.make(uri).svg_data_uri(scale=5))

    @app.post("/api/auth/mfa/enable")
    @limiter.limit("10/hour")
    @login_required
    def mfa_enable():
        secret = session.get("mfa_setup")
        if not (secret and security.verify_totp(secret, str((request.get_json(silent=True) or {}).get("code", "")))):
            abort(400, "That code did not match. Check the time on your phone and try again.")
        with engine.begin() as c:
            c.execute(update(db.users).where(db.users.c.id == session["uid"]).values(
                totp_secret=vault.encrypt(secret.encode()).decode(), totp_enabled=True))
        session.pop("mfa_setup")
        return jsonify(ok=True)

    @app.post("/api/auth/forgot")
    @limiter.limit("3/hour")
    @limiter.limit("3/hour", key_func=email_key)
    def forgot():  # same answer whether or not the email exists
        email = email_key().strip()
        with engine.connect() as c:
            user = c.execute(select(db.users).where(db.users.c.email == email)).first()
        if user:
            base = os.environ.get("FRONTEND_URL", allowed_origin)
            link = f"{base}/reset#token={security.issue_reset(engine, user.id)}"  # fragment: never sent to servers or logs

            def send():
                security.send_reset_email(user.email, link)

            send() if app.testing else threading.Thread(target=send, daemon=True).start()
        return jsonify(ok=True)

    @app.post("/api/auth/reset")
    @limiter.limit("10/hour")
    def reset():
        body = request.get_json(silent=True) or {}
        password = str(body.get("password", ""))[:256]
        if len(password) < 12:
            abort(400, "Password must be at least 12 characters")
        uid = security.consume_reset(engine, str(body.get("token", ""))[:200])
        if uid is None:
            abort(400, "This reset link is invalid or has expired")
        with engine.begin() as c:
            c.execute(update(db.users).where(db.users.c.id == uid).values(password_hash=generate_password_hash(password)))
        return jsonify(ok=True)

    @app.post("/api/auth/logout")
    def logout():
        session.clear()
        return jsonify(ok=True)

    @app.get("/api/auth/me")
    @login_required
    def me():
        user = user_by_id(session["uid"])
        return jsonify(email=user.email, mfa=bool(user.totp_enabled))

    @app.get("/api/meta")
    @login_required
    def meta():
        df = fetch(by_lga=False)
        with engine.connect() as c:
            names = sorted(r[0] for r in c.execute(select(db.lgas.c.name)))
        return jsonify(lgas=names, variables=VARS, start=df.period.min() if len(df) else "",
                       end=df.period.max() if len(df) else "")

    @app.get("/api/data")
    @login_required
    def data():
        return jsonify(nulls(fetch().drop(columns=["year", "month"])).to_dict("records"))

    @app.get("/api/summary")
    @login_required
    def summary():
        v = variable()
        s = fetch(by_lga=False).groupby("lga")[v].mean().round(2).sort_values(ascending=False)
        return jsonify([{"lga": k, "value": None if pd.isna(x) else x} for k, x in s.items()])

    @app.get("/api/boundaries")  # GeoJSON for maps; add ?variable=rain_mm&start=..&end=.. to attach a mean value
    @login_required
    def boundaries():
        vals = fetch(by_lga=False).groupby("lga")[variable()].mean().round(2) if "variable" in request.args else {}
        with engine.connect() as c:
            rows = c.execute(select(db.lgas)).all()
        feats = []
        for r in rows:
            val = vals.get(r.name) if len(vals) else None
            feats.append({"type": "Feature", "geometry": json.loads(r.geojson),
                          "properties": {"lga": r.name, "value": None if val is None or pd.isna(val) else val}})
        return jsonify(type="FeatureCollection", features=feats)

    def variables_arg(default=None, minimum=1):
        vs = list(dict.fromkeys(request.args.getlist("variable"))) or list(default or [])
        if len(vs) > 12 or len(vs) < minimum or any(v not in VARS for v in vs):
            abort(400, "Choose between %d and 12 known indicators" % minimum)
        return vs

    def choice(name, allowed, default):
        v = request.args.get(name, default)
        if v not in allowed:
            abort(400, f"Unknown {name}")
        return v

    def window(cols):
        """Monthly series for the chosen scope ('state' = mean of the selected LGAs, or one LGA), start/end honoured.
        Indicators with no data in the window are dropped so the matrix stays readable."""
        df = fetch(by_lga=False)
        scope = request.args.get("scope", "state")
        if scope != "state" and scope not in set(df["lga"]):
            abort(400, "Unknown scope")
        have = [c for c in cols if df[c].notna().any()]
        return analysis.wide(df, have, scope, request.args.getlist("lga") or None), have

    @app.get("/api/correlations")  # indicators against each other over time, optionally with a lag
    @login_required
    def correlations():
        cols = variables_arg(["rain_mm", "t2m", "t2m_max", "rh2m"], 2)
        method = choice("method", ("pearson", "spearman"), "pearson")
        deseason = request.args.get("deseasonalise", "1") != "0"
        w, have = window(cols)
        if len(have) < 2 or len(w) < analysis.MIN_PAIRS:
            return jsonify(variables=have, matrix=[], lag=[], months=len(w))
        base = analysis.deseasonalise(w, have) if deseason else w
        out = {"variables": have, "labels": {c: VARS[c] for c in have}, "months": len(w), "method": method,
               "deseasonalised": deseason, "start": w.index.min().strftime("%Y-%m"), "end": w.index.max().strftime("%Y-%m"),
               "matrix": analysis.matrix(base, have, method), "lag": []}
        la, lb = request.args.get("lag_a"), request.args.get("lag_b")
        if la in have and lb in have:
            try:
                ml = max(1, min(24, int(request.args.get("max_lag", 12))))
            except ValueError:
                abort(400, "max_lag must be a number")
            out["lag"] = analysis.lagged(base, la, lb, ml, method)
            out["lag_pair"] = [la, lb]
        return jsonify(out)

    @app.get("/api/trends")  # yearly trend for each indicator (complete years only)
    @login_required
    def trends():
        w, have = window(variables_arg(list(VARS)))
        return jsonify(variables=have, labels={c: VARS[c] for c in have}, trends=analysis.trends(w, have) if len(w) else {})

    def survey_frame(level=None):
        """Survey rows for ?indicator=&source=&year=&geo_level=, validated against what is in the database."""
        indicator, source = request.args.get("indicator", ""), request.args.get("source")
        geo = level or request.args.get("geo_level")
        s = db.surveys.c
        with engine.connect() as c:
            if not c.execute(select(s.indicator).where(s.indicator == indicator).limit(1)).first():
                abort(400, "Unknown indicator")
            q = select(db.surveys).where(s.indicator == indicator).order_by(s.survey_year, s.geo_name)
            if source:
                q = q.where(s.source == source)
            if geo:
                q = q.where(s.geo_level == geo)
            if request.args.get("year"):
                try:
                    q = q.where(s.survey_year == int(request.args["year"]))
                except ValueError:
                    abort(400, "year must be a number")
            df = pd.read_sql(q, c)
        if not source and df["source"].nunique() > 1:
            abort(400, "Several sources have this indicator; add ?source=")
        return df

    @app.get("/api/surveys/meta")
    @login_required
    def surveys_meta():
        s = db.surveys.c
        with engine.connect() as c:
            df = pd.read_sql(select(s.source, s.indicator, s.label, s.unit, s.survey_year, s.geo_level).distinct(), c)
        out = [{"source": src, "indicator": ind, "label": g["label"].iloc[0], "unit": g["unit"].iloc[0] or "",
                "years": sorted(int(y) for y in g["survey_year"].unique()), "geo_levels": sorted(g["geo_level"].unique())}
               for (src, ind), g in df.groupby(["source", "indicator"])]
        return jsonify(out)

    @app.get("/api/surveys")
    @login_required
    def surveys_data():
        return jsonify(nulls(survey_frame()).to_dict("records"))

    @app.get("/api/surveys/boundaries")  # GeoJSON with the LGA-level survey value attached, for the map
    @login_required
    def surveys_boundaries():
        df = survey_frame("lga")
        if df["survey_year"].nunique() > 1:
            df = df[df["survey_year"] == df["survey_year"].max()]  # latest round unless ?year= is given
        vals = df.set_index("geo_name")
        with engine.connect() as c:
            rows = c.execute(select(db.lgas)).all()
        feats = []
        for r in rows:
            row = vals.loc[r.name] if r.name in vals.index else None
            v = None if row is None or pd.isna(row["value"]) else float(row["value"])
            feats.append({"type": "Feature", "geometry": json.loads(r.geojson),
                          "properties": {"lga": r.name, "value": v, "note": None if row is None else row["note"],
                                         "n": None if row is None or pd.isna(row["n"]) else int(row["n"])}})
        return jsonify(type="FeatureCollection", features=feats)

    @app.get("/api/correlations/survey")  # across LGAs: survey value vs average climate over a window
    @login_required
    def survey_correlations():
        sv = survey_frame("lga")
        if sv.empty:
            return jsonify(results=[], n_lgas=0)
        year = int(sv["survey_year"].max() if not request.args.get("year") else request.args["year"])
        sv = sv[sv["survey_year"] == year].set_index("geo_name")["value"]
        start, end = request.args.get("start") or f"{year}-01", request.args.get("end") or f"{year}-12"
        method = choice("method", ("pearson", "spearman"), "pearson")
        cols = variables_arg(["rain_mm", "t2m", "t2m_max", "rh2m"])
        df = fetch(by_lga=False, start=start, end=end)
        means = df.groupby("lga")[cols].mean()
        res = [{"variable": c, "label": VARS[c], **analysis.cross_section(means[c], sv, method)}
               for c in cols if means[c].notna().any()]
        return jsonify(results=res, year=year, start=start, end=end, method=method, n_lgas=int(sv.notna().sum()),
                       note="Ecological comparison across LGAs: it shows where the two patterns coincide, "
                            "not that one causes the other.")

    @app.get("/api/export.xlsx")
    @login_required
    def export():
        df = fetch().drop(columns=["year", "month"])
        df["lga"] = df["lga"].map(safe_cell)
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as xw:
            df[["lga", "period", *VARS]].rename(columns={"lga": "LGA", "period": "Month", **VARS}).to_excel(
                xw, sheet_name="Data", index=False)
            for v, label in VARS.items():  # one LGA x month grid per indicator, ready for mapping/modelling
                grid = df.pivot(index="lga", columns="period", values=v)
                grid.to_excel(xw, sheet_name=label[:28].replace("/", "-"), index_label="LGA")
            with engine.connect() as c:
                sv = pd.read_sql(select(db.surveys), c)
            if len(sv):
                for col in ("geo_name", "label", "note"):
                    sv[col] = sv[col].map(safe_cell)
                sv.rename(columns={"survey_year": "year"}).to_excel(xw, sheet_name="Survey indicators", index=False)
            pd.DataFrame(SOURCES, columns=["Item", "Source / note"]).to_excel(xw, sheet_name="About", index=False)
            for ws in xw.book.worksheets:
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="0F766E")
                ws.freeze_panes = "A2" if ws.title == "About" else "B2"
                for i, col in enumerate(ws.columns, 1):
                    width = max(len(str(c.value or "")) for c in col) + 2
                    ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width))
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name="ondo_climate_lga_monthly.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    return app


if __name__ == "__main__":
    create_app().run(port=5000)
